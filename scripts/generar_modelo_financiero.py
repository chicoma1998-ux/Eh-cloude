# -*- coding: utf-8 -*-
"""Construye el modelo financiero inmobiliario: flujo de caja, cronograma
valorizado, estado de resultados e indicadores (VAN/TIR)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

OUT = "/home/user/Eh-cloude/entregables/Modelo_Financiero_Proyecto_Inmobiliario.xlsx"

# ---------- Estilos ----------
ARIAL = "Arial"
FONT = Font(name=ARIAL, size=10, color="000000")
FONT_BLUE = Font(name=ARIAL, size=10, color="0000FF")       # inputs
FONT_GREEN = Font(name=ARIAL, size=10, color="008000")      # links a otra hoja
FONT_BOLD = Font(name=ARIAL, size=10, bold=True, color="000000")
FONT_TITLE = Font(name=ARIAL, size=14, bold=True, color="FFFFFF")
FONT_SECTION = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
FONT_HDR = Font(name=ARIAL, size=9, bold=True, color="FFFFFF")

FILL_TITLE = PatternFill("solid", fgColor="1F3864")   # azul oscuro
FILL_SECTION = PatternFill("solid", fgColor="2E5496")  # azul medio
FILL_HDR = PatternFill("solid", fgColor="305496")
FILL_YELLOW = PatternFill("solid", fgColor="FFF2CC")   # celdas a completar
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
FILL_ALT = PatternFill("solid", fgColor="F2F2F2")

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

MONEY = '#,##0;(#,##0);-'
MONEY2 = '"S/ "#,##0;("S/ "#,##0);-'
PCT = '0.0%'
PCT2 = '0.00%'
NUM = '#,##0'

N = 24                     # meses
COL0 = 3                   # primera columna de meses = C
LAST = COL0 + N - 1        # Z (26)
TOTAL_COL = 2              # columna B = Total

def cl(i):  # column letter para mes i (1..N)
    return get_column_letter(COL0 + i - 1)

def month_cols():
    return [get_column_letter(c) for c in range(COL0, LAST + 1)]

wb = openpyxl.Workbook()
# Forzar recálculo completo al abrir (Excel/Sheets muestran los valores al instante)
wb.calculation.fullCalcOnLoad = True

# =====================================================================
#  HOJA: SUPUESTOS
# =====================================================================
ws = wb.active
ws.title = "Supuestos"
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 46
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 42

def title(ws, text, span):
    ws.merge_cells(f"A1:{get_column_letter(span)}1")
    c = ws["A1"]; c.value = text; c.font = FONT_TITLE; c.fill = FILL_TITLE; c.alignment = CENTER
    ws.row_dimensions[1].height = 26

title(ws, "SUPUESTOS Y PARÁMETROS DEL PROYECTO", 3)

S = {}  # nombre -> fila
def put(row, label, value, fmt=None, blue=False, note=None, bold=False, section=False):
    a = ws.cell(row=row, column=1, value=label)
    b = ws.cell(row=row, column=2, value=value)
    if section:
        ws.merge_cells(f"A{row}:C{row}")
        a.font = FONT_SECTION; a.fill = FILL_SECTION; a.alignment = LEFT
        return
    a.font = FONT_BOLD if bold else FONT
    a.alignment = LEFT
    b.alignment = RIGHT
    if blue:
        b.font = FONT_BLUE; b.fill = FILL_INPUT
    else:
        b.font = FONT
    if fmt:
        b.number_format = fmt
    if note:
        c = ws.cell(row=row, column=3, value=note)
        c.font = Font(name=ARIAL, size=9, italic=True, color="808080")
        c.alignment = LEFT

r = 3
put(r, "1. DATOS GENERALES", None, section=True); r += 1
put(r, "Nombre del proyecto", "VILLA EL MILAGRO (AVN – Techo Propio)", blue=True, note="Editable"); S['proy']=r; r += 1
put(r, "Ubicación", "[Distrito, Provincia]", blue=True, note="Completar"); r += 1
put(r, "Horizonte de evaluación (meses)", N, NUM, blue=True); S['meses']=r; r += 1
put(r, "Área techada vendible (m²)", 3500, NUM, blue=True); S['area']=r; r += 1
put(r, "N° de unidades inmobiliarias", 20, NUM, blue=True); S['unid']=r; r += 1

r += 1
put(r, "2. INGRESOS Y VENTAS", None, section=True); r += 1
put(r, "Precio de venta promedio por unidad (S/)", 560000, MONEY, blue=True); S['precio']=r; r += 1
put(r, "Ventas totales del proyecto (S/)", f"=B{S['unid']}*B{S['precio']}", MONEY); S['ventas']=r; r += 1
put(r, "Cuota inicial de clientes (%)", 0.10, PCT, blue=True,
    note="% que paga el cliente durante la obra"); S['ini']=r; r += 1
put(r, "Saldo financiado por crédito hipotecario (%)", f"=1-B{S['ini']}", PCT); S['saldo']=r; r += 1

r += 1
put(r, "3. COSTOS DEL PROYECTO", None, section=True); r += 1
put(r, "Costo del terreno (S/)", 2500000, MONEY, blue=True); S['terreno']=r; r += 1
put(r, "Costo de construcción por m² (S/)", 1650, MONEY, blue=True); S['cm2']=r; r += 1
put(r, "Costo total de construcción (S/)", f"=B{S['area']}*B{S['cm2']}", MONEY); S['constr']=r; r += 1
put(r, "Licencias, proyectos y estudios (% de construcción)", 0.06, PCT, blue=True); S['plic']=r; r += 1
put(r, "Monto licencias/proyectos (S/)", f"=B{S['constr']}*B{S['plic']}", MONEY); S['lic']=r; r += 1
put(r, "Gastos de comercialización (% de ventas)", 0.03, PCT, blue=True); S['pcom']=r; r += 1
put(r, "Monto comercialización (S/)", f"=B{S['ventas']}*B{S['pcom']}", MONEY); S['com']=r; r += 1
put(r, "Gastos administrativos / gerencia (% de ventas)", 0.04, PCT, blue=True); S['padm']=r; r += 1
put(r, "Monto administrativos (S/)", f"=B{S['ventas']}*B{S['padm']}", MONEY); S['adm']=r; r += 1

r += 1
put(r, "4. FINANCIAMIENTO / GARANTÍA (CARTA FIANZA)", None, section=True); r += 1
put(r, "Línea de financiamiento (% de la construcción)", 0.60, PCT, blue=True,
    note="Porcentaje de la obra respaldado por financiamiento / carta fianza"); S['pfin']=r; r += 1
put(r, "Monto máximo de la línea (S/)", f"=B{S['constr']}*B{S['pfin']}", MONEY); S['linea']=r; r += 1
put(r, "Tasa de interés anual (TEA)", 0.12, PCT2, blue=True); S['tea']=r; r += 1
put(r, "Tasa de interés mensual (TEM)", f"=(1+B{S['tea']})^(1/12)-1", PCT2); S['tem']=r; r += 1

r += 1
put(r, "5. EVALUACIÓN ECONÓMICA", None, section=True); r += 1
put(r, "Tasa de descuento anual (COK)", 0.15, PCT2, blue=True); S['cok']=r; r += 1
put(r, "Tasa de descuento mensual", f"=(1+B{S['cok']})^(1/12)-1", PCT2); S['cokm']=r; r += 1
put(r, "Impuesto a la renta (%)", 0.295, PCT2, blue=True); S['ir']=r; r += 1

r += 2
leg = ws.cell(row=r, column=1, value="Leyenda:  Celdas amarillas = datos editables (inputs).  Celdas blancas = calculadas por fórmula.")
ws.merge_cells(f"A{r}:C{r}")
leg.font = Font(name=ARIAL, size=9, italic=True, color="C00000")

def SUP(key):
    return f"Supuestos!$B${S[key]}"

# =====================================================================
#  HOJA: CRONOGRAMA VALORIZADO
# =====================================================================
wc = wb.create_sheet("Cronograma_Valorizado")
wc.sheet_view.showGridLines = False
wc.column_dimensions['A'].width = 40
wc.column_dimensions['B'].width = 15
for c in range(COL0, LAST + 1):
    wc.column_dimensions[get_column_letter(c)].width = 11

wc.merge_cells(f"A1:{get_column_letter(LAST)}1")
t = wc["A1"]; t.value = "CRONOGRAMA VALORIZADO DE OBRA Y CALENDARIO DE DESEMBOLSOS"
t.font = FONT_TITLE; t.fill = FILL_TITLE; t.alignment = CENTER
wc.row_dimensions[1].height = 26

# encabezado meses (fila 3)
hr = 3
wc.cell(row=hr, column=1, value="Concepto").font = FONT_HDR
wc.cell(row=hr, column=1).fill = FILL_HDR; wc.cell(row=hr, column=1).alignment = LEFT
wc.cell(row=hr, column=TOTAL_COL, value="Total").font = FONT_HDR
wc.cell(row=hr, column=TOTAL_COL).fill = FILL_HDR; wc.cell(row=hr, column=TOTAL_COL).alignment = CENTER
for i in range(1, N + 1):
    c = wc.cell(row=hr, column=COL0 + i - 1, value=f"Mes {i}")
    c.font = FONT_HDR; c.fill = FILL_HDR; c.alignment = CENTER

# perfil S-curve (pesos) meses 4..21
weights = {i: 0 for i in range(1, N + 1)}
scurve = {4:2,5:3,6:4,7:5,8:6,9:7,10:8,11:8,12:8,13:8,14:8,15:7,16:7,17:6,18:5,19:4,20:3,21:2}
weights.update(scurve)

CR = {}
rr = 4
# fila pesos
wc.cell(row=rr, column=1, value="Perfil de avance de obra (peso S-curve)").font = FONT
wc.cell(row=rr, column=1).alignment = LEFT
wc.cell(row=rr, column=TOTAL_COL, value=f"=SUM(C{rr}:{get_column_letter(LAST)}{rr})").font = FONT_BOLD
wc.cell(row=rr, column=TOTAL_COL).number_format = NUM
for i in range(1, N + 1):
    c = wc.cell(row=rr, column=COL0 + i - 1, value=weights[i])
    c.font = FONT_BLUE; c.fill = FILL_INPUT; c.number_format = NUM; c.alignment = CENTER
CR['peso'] = rr; rr += 1

# % avance del mes
wc.cell(row=rr, column=1, value="% de avance del mes").font = FONT
wc.cell(row=rr, column=TOTAL_COL, value=f"=SUM(C{rr}:{get_column_letter(LAST)}{rr})").font = FONT_BOLD
wc.cell(row=rr, column=TOTAL_COL).number_format = PCT
for i in range(1, N + 1):
    col = cl(i)
    c = wc.cell(row=rr, column=COL0 + i - 1,
                value=f"={col}{CR['peso']}/$B${CR['peso']}")
    c.number_format = PCT; c.font = FONT
CR['pct'] = rr; rr += 1

# valorización del mes
wc.cell(row=rr, column=1, value="Valorización de obra del mes (S/)").font = FONT_BOLD
wc.cell(row=rr, column=TOTAL_COL, value=f"=SUM(C{rr}:{get_column_letter(LAST)}{rr})").font = FONT_BOLD
wc.cell(row=rr, column=TOTAL_COL).number_format = MONEY
for i in range(1, N + 1):
    col = cl(i)
    c = wc.cell(row=rr, column=COL0 + i - 1, value=f"={col}{CR['pct']}*{SUP('constr')}")
    c.number_format = MONEY; c.font = FONT
CR['val'] = rr; rr += 1

# valorización acumulada
wc.cell(row=rr, column=1, value="Valorización acumulada (S/)").font = FONT
wc.cell(row=rr, column=TOTAL_COL).value = None
for i in range(1, N + 1):
    col = cl(i)
    if i == 1:
        f = f"={col}{CR['val']}"
    else:
        f = f"={get_column_letter(COL0+i-2)}{rr}+{col}{CR['val']}"
    c = wc.cell(row=rr, column=COL0 + i - 1, value=f)
    c.number_format = MONEY; c.font = FONT
CR['valac'] = rr; rr += 1

# % avance acumulado
wc.cell(row=rr, column=1, value="% avance acumulado").font = FONT
for i in range(1, N + 1):
    col = cl(i)
    c = wc.cell(row=rr, column=COL0 + i - 1, value=f"={col}{CR['valac']}/{SUP('constr')}")
    c.number_format = PCT; c.font = FONT
CR['pctac'] = rr; rr += 1

rr += 1
# desembolso bancario del mes
wc.cell(row=rr, column=1, value="Desembolso de financiamiento del mes (S/)").font = FONT_BOLD
wc.cell(row=rr, column=1).fill = FILL_TOTAL
wc.cell(row=rr, column=TOTAL_COL, value=f"=SUM(C{rr}:{get_column_letter(LAST)}{rr})").font = FONT_BOLD
wc.cell(row=rr, column=TOTAL_COL).number_format = MONEY
wc.cell(row=rr, column=TOTAL_COL).fill = FILL_TOTAL
for i in range(1, N + 1):
    col = cl(i)
    c = wc.cell(row=rr, column=COL0 + i - 1, value=f"={col}{CR['val']}*{SUP('pfin')}")
    c.number_format = MONEY; c.font = FONT; c.fill = FILL_TOTAL
CR['des'] = rr; rr += 1

# desembolso acumulado
wc.cell(row=rr, column=1, value="Desembolso de financiamiento acumulado (S/)").font = FONT
for i in range(1, N + 1):
    col = cl(i)
    if i == 1:
        f = f"={col}{CR['des']}"
    else:
        f = f"={get_column_letter(COL0+i-2)}{rr}+{col}{CR['des']}"
    c = wc.cell(row=rr, column=COL0 + i - 1, value=f)
    c.number_format = MONEY; c.font = FONT
CR['desac'] = rr; rr += 1

wc.cell(row=CR['des'], column=COL0).comment = Comment(
    "Desembolso = % de la obra respaldado por financiamiento/carta fianza × valorización del mes. "
    "Sustenta el calendario de desembolsos y el dimensionamiento de la carta fianza.", "Modelo")

nota_r = rr + 1
ncell = wc.cell(row=nota_r, column=1,
    value="Nota: el desembolso bancario se ata a la valorización mensual de obra (avance físico real). "
          "Ajuste el perfil S-curve (fila azul) según el cronograma de ejecución del proyecto.")
wc.merge_cells(f"A{nota_r}:{get_column_letter(LAST)}{nota_r}")
ncell.font = Font(name=ARIAL, size=9, italic=True, color="808080")

def CRO(key, i):
    return f"Cronograma_Valorizado!{cl(i)}${CR[key]}"

# =====================================================================
#  HOJA: FLUJO DE CAJA
# =====================================================================
wf = wb.create_sheet("Flujo_Caja")
wf.sheet_view.showGridLines = False
wf.column_dimensions['A'].width = 42
wf.column_dimensions['B'].width = 15
for c in range(COL0, LAST + 1):
    wf.column_dimensions[get_column_letter(c)].width = 11

wf.merge_cells(f"A1:{get_column_letter(LAST)}1")
t = wf["A1"]; t.value = "FLUJO DE CAJA DEL PROYECTO (mensual, en Soles)"
t.font = FONT_TITLE; t.fill = FILL_TITLE; t.alignment = CENTER
wf.row_dimensions[1].height = 26

hr = 3
wf.cell(row=hr, column=1, value="Concepto").font = FONT_HDR
wf.cell(row=hr, column=1).fill = FILL_HDR; wf.cell(row=hr, column=1).alignment = LEFT
wf.cell(row=hr, column=TOTAL_COL, value="Total").font = FONT_HDR
wf.cell(row=hr, column=TOTAL_COL).fill = FILL_HDR; wf.cell(row=hr, column=TOTAL_COL).alignment = CENTER
for i in range(1, N + 1):
    c = wf.cell(row=hr, column=COL0 + i - 1, value=f"Mes {i}")
    c.font = FONT_HDR; c.fill = FILL_HDR; c.alignment = CENTER

F = {}
def frow(row, label, section=False, total=False, bold=False, fill=None):
    a = wf.cell(row=row, column=1, value=label)
    if section:
        wf.merge_cells(f"A{row}:{get_column_letter(LAST)}{row}")
        a.font = FONT_SECTION; a.fill = FILL_SECTION; a.alignment = LEFT
        return
    a.font = FONT_BOLD if (bold or total) else FONT
    a.alignment = LEFT
    if fill:
        a.fill = fill

def fillrow(row, func, fmt=MONEY, total_formula=True, font=FONT, fill=None):
    for i in range(1, N + 1):
        c = wf.cell(row=row, column=COL0 + i - 1, value=func(i))
        c.number_format = fmt; c.font = font
        if fill: c.fill = fill
    if total_formula:
        c = wf.cell(row=row, column=TOTAL_COL,
                    value=f"=SUM(C{row}:{get_column_letter(LAST)}{row})")
        c.number_format = fmt; c.font = FONT_BOLD
        if fill: c.fill = fill

rr = 4
# --- drivers ---
frow(rr, "DISTRIBUCIÓN TEMPORAL (perfiles editables)", section=True); rr += 1

# perfil recaudo cuota inicial (meses 2..20)
frow(rr, "Perfil de recaudo de cuota inicial (peso)")
recw = {i: (1 if 2 <= i <= 20 else 0) for i in range(1, N + 1)}
for i in range(1, N + 1):
    c = wf.cell(row=rr, column=COL0 + i - 1, value=recw[i])
    c.font = FONT_BLUE; c.fill = FILL_INPUT; c.number_format = NUM; c.alignment = CENTER
wf.cell(row=rr, column=TOTAL_COL, value=f"=SUM(C{rr}:{get_column_letter(LAST)}{rr})").font = FONT_BOLD
wf.cell(row=rr, column=TOTAL_COL).number_format = NUM
F['recw'] = rr; rr += 1

# perfil entrega/cobro saldo (meses 22..24)
frow(rr, "Perfil de entrega / cobro de saldo (peso)")
entw = {i: 0 for i in range(1, N + 1)}
entw.update({22:1, 23:2, 24:2})
for i in range(1, N + 1):
    c = wf.cell(row=rr, column=COL0 + i - 1, value=entw[i])
    c.font = FONT_BLUE; c.fill = FILL_INPUT; c.number_format = NUM; c.alignment = CENTER
wf.cell(row=rr, column=TOTAL_COL, value=f"=SUM(C{rr}:{get_column_letter(LAST)}{rr})").font = FONT_BOLD
wf.cell(row=rr, column=TOTAL_COL).number_format = NUM
F['entw'] = rr; rr += 1

rr += 1
# --- INGRESOS ---
frow(rr, "A. INGRESOS", section=True); rr += 1

frow(rr, "Cuota inicial de clientes")
def f_ini(i):
    col = cl(i)
    return (f"=({col}${F['recw']}/$B${F['recw']})*{SUP('ventas')}*{SUP('ini')}")
fillrow(rr, f_ini)
F['ini'] = rr; rr += 1

frow(rr, "Cobro de saldo (créditos hipotecarios)")
def f_saldo(i):
    col = cl(i)
    return (f"=({col}${F['entw']}/$B${F['entw']})*{SUP('ventas')}*{SUP('saldo')}")
fillrow(rr, f_saldo)
F['saldo'] = rr; rr += 1

frow(rr, "Desembolsos de financiamiento")
fillrow(rr, lambda i: f"={CRO('des', i)}", font=FONT_GREEN)
F['des'] = rr; rr += 1

frow(rr, "Total ingresos", total=True)
fillrow(rr, lambda i: f"=SUM({cl(i)}{F['ini']}:{cl(i)}{F['des']})", fill=FILL_TOTAL, font=FONT_BOLD)
F['ting'] = rr; rr += 1

rr += 1
# --- EGRESOS ---
frow(rr, "B. EGRESOS OPERATIVOS", section=True); rr += 1

frow(rr, "Compra de terreno")
fillrow(rr, lambda i: f"=-{SUP('terreno')}" if i == 1 else 0)
F['terr'] = rr; rr += 1

frow(rr, "Costo de construcción (valorización de obra)")
fillrow(rr, lambda i: f"=-{CRO('val', i)}", font=FONT_GREEN)
F['con'] = rr; rr += 1

frow(rr, "Licencias, proyectos y estudios")
fillrow(rr, lambda i: f"=-{SUP('lic')}/3" if i in (1, 2, 3) else 0)
F['lic'] = rr; rr += 1

frow(rr, "Gastos de comercialización")
fillrow(rr, lambda i: f"=-{SUP('com')}/20" if 3 <= i <= 22 else 0)
F['com'] = rr; rr += 1

frow(rr, "Gastos administrativos / gerencia")
fillrow(rr, lambda i: f"=-{SUP('adm')}/{N}")
F['adm'] = rr; rr += 1

frow(rr, "Total egresos operativos", total=True)
fillrow(rr, lambda i: f"=SUM({cl(i)}{F['terr']}:{cl(i)}{F['adm']})", fill=FILL_TOTAL, font=FONT_BOLD)
F['teg'] = rr; rr += 1

rr += 1
# --- FINANCIAMIENTO / DEUDA ---
frow(rr, "C. SERVICIO DE LA DEUDA / FINANCIAMIENTO", section=True); rr += 1

frow(rr, "Saldo inicial de la línea")
F['sini'] = rr
# saldo final placeholder row known after; define order: sini, desem, amort, int, sfin
F['desL'] = rr + 1
F['amort'] = rr + 2
F['intr'] = rr + 3
F['sfin'] = rr + 4

for i in range(1, N + 1):
    col = cl(i)
    if i == 1:
        f = 0
    else:
        f = f"={get_column_letter(COL0+i-2)}{F['sfin']}"
    c = wf.cell(row=F['sini'], column=COL0 + i - 1, value=f)
    c.number_format = MONEY; c.font = FONT
rr += 1

frow(rr, "(+) Desembolso del mes")
fillrow(rr, lambda i: f"={CRO('des', i)}", font=FONT_GREEN)
rr += 1

frow(rr, "(-) Amortización del mes")
def f_amort(i):
    col = cl(i)
    return f"=-({col}${F['entw']}/$B${F['entw']})*{CRO('des','TOT') if False else ''}Cronograma_Valorizado!$B${CR['des']}"
# amortización = -% entrega * total desembolsado
def f_amort2(i):
    col = cl(i)
    return f"=-({col}${F['entw']}/$B${F['entw']})*Cronograma_Valorizado!$B${CR['des']}"
fillrow(rr, f_amort2)
rr += 1

frow(rr, "(-) Intereses del mes (gasto financiero)")
def f_int(i):
    col = cl(i)
    return f"=-{col}{F['sini']}*{SUP('tem')}"
fillrow(rr, f_int)
rr += 1

frow(rr, "Saldo final de la línea", bold=True)
for i in range(1, N + 1):
    col = cl(i)
    f = f"={col}{F['sini']}+{col}{F['desL']}+{col}{F['amort']}"
    c = wf.cell(row=F['sfin'], column=COL0 + i - 1, value=f)
    c.number_format = MONEY; c.font = FONT_BOLD; c.fill = FILL_TOTAL
wf.cell(row=F['sfin'], column=1).fill = FILL_TOTAL
rr += 1

rr += 1
# --- FLUJOS DE CAJA ---
frow(rr, "D. FLUJOS DE CAJA", section=True); rr += 1

frow(rr, "Flujo económico del mes (sin financiamiento)", bold=True)
def f_eco(i):
    col = cl(i)
    return f"={col}{F['ini']}+{col}{F['saldo']}+{col}{F['teg']}"
fillrow(rr, f_eco, fill=FILL_TOTAL, font=FONT_BOLD)
F['eco'] = rr; rr += 1

frow(rr, "Flujo económico acumulado")
for i in range(1, N + 1):
    col = cl(i)
    f = f"={col}{F['eco']}" if i == 1 else f"={get_column_letter(COL0+i-2)}{rr}+{col}{F['eco']}"
    c = wf.cell(row=rr, column=COL0 + i - 1, value=f)
    c.number_format = MONEY; c.font = FONT
F['ecoac'] = rr; rr += 1

frow(rr, "Flujo financiero del mes (con financiamiento)", bold=True)
def f_fin(i):
    col = cl(i)
    return f"={col}{F['eco']}+{col}{F['desL']}+{col}{F['amort']}+{col}{F['intr']}"
fillrow(rr, f_fin, fill=FILL_TOTAL, font=FONT_BOLD)
F['fin'] = rr; rr += 1

frow(rr, "Flujo financiero acumulado (caja del proyecto)", bold=True)
for i in range(1, N + 1):
    col = cl(i)
    f = f"={col}{F['fin']}" if i == 1 else f"={get_column_letter(COL0+i-2)}{rr}+{col}{F['fin']}"
    c = wf.cell(row=rr, column=COL0 + i - 1, value=f)
    c.number_format = MONEY; c.font = FONT_BOLD
    c.fill = FILL_TOTAL
F['finac'] = rr; rr += 1

def FLU(key, i):
    return f"Flujo_Caja!{cl(i)}${F[key]}"
def FLU_TOT(key):
    return f"Flujo_Caja!$B${F[key]}"

# =====================================================================
#  HOJA: ESTADO DE RESULTADOS
# =====================================================================
we = wb.create_sheet("Estado_Resultados")
we.sheet_view.showGridLines = False
we.column_dimensions['A'].width = 46
we.column_dimensions['B'].width = 18
we.column_dimensions['C'].width = 14

we.merge_cells("A1:C1")
t = we["A1"]; t.value = "ESTADO DE RESULTADOS PROYECTADO DEL PROYECTO"
t.font = FONT_TITLE; t.fill = FILL_TITLE; t.alignment = CENTER
we.row_dimensions[1].height = 26

E = {}
def erow(row, label, value=None, fmt=MONEY, bold=False, section=False, pct=False, fill=None):
    a = we.cell(row=row, column=1, value=label)
    b = we.cell(row=row, column=2, value=value)
    if section:
        we.merge_cells(f"A{row}:C{row}")
        a.font = FONT_SECTION; a.fill = FILL_SECTION; a.alignment = LEFT
        return
    a.font = FONT_BOLD if bold else FONT; a.alignment = LEFT
    b.font = FONT_BOLD if bold else FONT; b.alignment = RIGHT
    b.number_format = PCT if pct else fmt
    if fill:
        a.fill = fill; b.fill = fill

r = 3
erow(r, "INGRESOS", section=True); r += 1
erow(r, "Ventas totales del proyecto", f"={SUP('ventas')}", bold=True); E['ventas']=r; r += 1
r += 1
erow(r, "COSTO DE VENTAS", section=True); r += 1
erow(r, "(-) Costo del terreno", f"=-{SUP('terreno')}"); E['terr']=r; r += 1
erow(r, "(-) Costo de construcción", f"=-{SUP('constr')}"); E['con']=r; r += 1
erow(r, "(-) Licencias, proyectos y estudios", f"=-{SUP('lic')}"); E['lic']=r; r += 1
erow(r, "UTILIDAD BRUTA", f"=B{E['ventas']}+B{E['terr']}+B{E['con']}+B{E['lic']}", bold=True, fill=FILL_TOTAL); E['bruta']=r; r += 1
erow(r, "Margen bruto", f"=B{E['bruta']}/B{E['ventas']}", pct=True); E['mbruta']=r; r += 1
r += 1
erow(r, "GASTOS OPERATIVOS", section=True); r += 1
erow(r, "(-) Gastos de comercialización", f"=-{SUP('com')}"); E['com']=r; r += 1
erow(r, "(-) Gastos administrativos / gerencia", f"=-{SUP('adm')}"); E['adm']=r; r += 1
erow(r, "UTILIDAD OPERATIVA", f"=B{E['bruta']}+B{E['com']}+B{E['adm']}", bold=True, fill=FILL_TOTAL); E['oper']=r; r += 1
r += 1
erow(r, "RESULTADO FINANCIERO E IMPUESTOS", section=True); r += 1
erow(r, "(-) Gastos financieros (intereses de la línea)", f"={FLU_TOT('intr')}"); E['fin']=r; r += 1
erow(r, "UTILIDAD ANTES DE IMPUESTOS", f"=B{E['oper']}+B{E['fin']}", bold=True, fill=FILL_TOTAL); E['uai']=r; r += 1
erow(r, "(-) Impuesto a la renta", f"=-MAX(B{E['uai']},0)*{SUP('ir')}"); E['ir']=r; r += 1
erow(r, "UTILIDAD NETA", f"=B{E['uai']}+B{E['ir']}", bold=True, fill=FILL_TOTAL); E['neta']=r; r += 1
erow(r, "Margen neto", f"=B{E['neta']}/B{E['ventas']}", pct=True); E['mneta']=r; r += 1
r += 2
note = we.cell(row=r, column=1, value="Nota: cifras en valores netos (sin IGV). El intereses proviene del flujo de caja financiero.")
we.merge_cells(f"A{r}:C{r}")
note.font = Font(name=ARIAL, size=9, italic=True, color="808080")

# =====================================================================
#  HOJA: RESUMEN / INDICADORES
# =====================================================================
wr = wb.create_sheet("Resumen")
wr.sheet_view.showGridLines = False
wr.column_dimensions['A'].width = 48
wr.column_dimensions['B'].width = 20

wr.merge_cells("A1:B1")
t = wr["A1"]; t.value = "RESUMEN EJECUTIVO E INDICADORES"
t.font = FONT_TITLE; t.fill = FILL_TITLE; t.alignment = CENTER
wr.row_dimensions[1].height = 26

def rrow(row, label, value=None, fmt=MONEY, bold=False, section=False, pct=False):
    a = wr.cell(row=row, column=1, value=label)
    b = wr.cell(row=row, column=2, value=value)
    if section:
        wr.merge_cells(f"A{row}:B{row}")
        a.font = FONT_SECTION; a.fill = FILL_SECTION; a.alignment = LEFT
        return
    a.font = FONT_BOLD if bold else FONT; a.alignment = LEFT
    b.font = FONT_BOLD if bold else FONT; b.alignment = RIGHT
    b.number_format = PCT2 if pct else fmt
    if bold:
        a.fill = FILL_TOTAL; b.fill = FILL_TOTAL

# rango de flujos para VAN/TIR
eco_range = f"Flujo_Caja!{cl(1)}{F['eco']}:{cl(N)}{F['eco']}"
fin_range = f"Flujo_Caja!{cl(1)}{F['fin']}:{cl(N)}{F['fin']}"

r = 3
rrow(r, "INVERSIÓN Y ESTRUCTURA", section=True); r += 1
rrow(r, "Ventas totales (S/)", f"={SUP('ventas')}"); r += 1
rrow(r, "Inversión total del proyecto (S/)", f"={SUP('terreno')}+{SUP('constr')}+{SUP('lic')}+{SUP('com')}+{SUP('adm')}"); Rinv=r; r += 1
rrow(r, "Monto máximo de financiamiento / carta fianza (S/)", f"={SUP('linea')}"); r += 1
rrow(r, "Apalancamiento (financiamiento / inversión)", f"={SUP('linea')}/B{Rinv}", pct=True); r += 1
rrow(r, "Aporte propio del desarrollador (S/)", f"=B{Rinv}-{SUP('linea')}"); r += 1
r += 1
rrow(r, "RENTABILIDAD", section=True); r += 1
rrow(r, "Utilidad neta (S/)", f"=Estado_Resultados!B{E['neta']}", bold=True); r += 1
rrow(r, "Margen neto", f"=Estado_Resultados!B{E['mneta']}", pct=True); r += 1
rrow(r, "ROI sobre inversión total", f"=Estado_Resultados!B{E['neta']}/B{Rinv}", pct=True); r += 1
r += 1
rrow(r, "EVALUACIÓN ECONÓMICA (sin financiamiento)", section=True); r += 1
rrow(r, "VAN económico (S/)", f"=NPV({SUP('cokm')},{eco_range})", bold=True); r += 1
rrow(r, "TIR económica mensual", f"=IRR({eco_range})", pct=True); Rte=r; r += 1
rrow(r, "TIR económica anual", f"=(1+B{Rte})^12-1", pct=True); r += 1
r += 1
rrow(r, "EVALUACIÓN FINANCIERA (con financiamiento)", section=True); r += 1
rrow(r, "VAN financiero (S/)", f"=NPV({SUP('cokm')},{fin_range})", bold=True); r += 1
rrow(r, "TIR financiera mensual", f"=IRR({fin_range})", pct=True); Rtf=r; r += 1
rrow(r, "TIR financiera anual", f"=(1+B{Rtf})^12-1", pct=True); r += 1
r += 2
note = wr.cell(row=r, column=1,
    value="El VAN se calcula con la tasa de descuento mensual (COK) de la hoja Supuestos. "
          "La TIR se anualiza como (1+TIR mensual)^12-1.")
wr.merge_cells(f"A{r}:B{r}")
note.font = Font(name=ARIAL, size=9, italic=True, color="808080")

# =====================================================================
#  HOJA: PORTADA (al inicio)
# =====================================================================
wp = wb.create_sheet("Portada", 0)
wp.sheet_view.showGridLines = False
wp.column_dimensions['A'].width = 4
wp.column_dimensions['B'].width = 60
wp.column_dimensions['C'].width = 40
wp.merge_cells("B2:C2")
c = wp["B2"]; c.value = "MODELO FINANCIERO – PROYECTO INMOBILIARIO"
c.font = Font(name=ARIAL, size=18, bold=True, color="1F3864"); c.alignment = LEFT
wp.merge_cells("B3:C3")
c = wp["B3"]; c.value = "Flujo de caja · Cronograma valorizado · Estado de resultados · Indicadores"
c.font = Font(name=ARIAL, size=11, italic=True, color="808080")

info = [
    ("Proyecto:", "VILLA EL MILAGRO – AVN, Programa Techo Propio"),
    ("Empresa desarrolladora:", "INVERSIONES ESTRATÉGICAS INTEGRALES S.A.C."),
    ("Elaborado por:", "Enrique Huidobro – Economista (RUC 10723712039)"),
    ("Fecha:", "31/07/2026"),
    ("Moneda:", "Soles (S/)"),
    ("Destino:", "Sustento para carta fianza ante empresas aseguradoras"),
]
rr = 5
for lab, val in info:
    a = wp.cell(row=rr, column=2, value=lab); a.font = FONT_BOLD; a.alignment = LEFT
    b = wp.cell(row=rr, column=3, value=val); b.font = FONT_BLUE; b.fill = FILL_INPUT; b.alignment = LEFT
    rr += 1

rr += 1
wp.merge_cells(f"B{rr}:C{rr}")
c = wp.cell(row=rr, column=2, value="CONTENIDO DEL LIBRO"); c.font = FONT_SECTION; c.fill = FILL_SECTION
rr += 1
contenido = [
    ("1. Supuestos", "Parámetros editables del proyecto (celdas amarillas)."),
    ("2. Cronograma_Valorizado", "Avance de obra y desembolsos bancarios por mes."),
    ("3. Flujo_Caja", "Ingresos, egresos y flujo neto mensual (24 meses)."),
    ("4. Estado_Resultados", "P&L proyectado, utilidad y márgenes."),
    ("5. Resumen", "Indicadores: VAN, TIR, ROI, apalancamiento."),
]
for lab, desc in contenido:
    a = wp.cell(row=rr, column=2, value=lab); a.font = FONT_BOLD; a.alignment = LEFT
    b = wp.cell(row=rr, column=3, value=desc); b.font = FONT; b.alignment = LEFT
    rr += 1

rr += 1
wp.merge_cells(f"B{rr}:C{rr}")
c = wp.cell(row=rr, column=2,
    value="Cómo usar: edite solo las celdas amarillas (inputs). Todo lo demás se recalcula automáticamente.")
c.font = Font(name=ARIAL, size=9, italic=True, color="C00000")

wb.save(OUT)
print("Guardado:", OUT)
